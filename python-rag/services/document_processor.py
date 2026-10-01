import io
from pathlib import Path
from typing import List, Dict, Any
import pypdf
import docx

class DocumentProcessor:
    """
    Extracts text from PDF, DOCX, TXT, and Markdown documents.
    Tracks page numbers or section indices for precise citation generation.
    """

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    @classmethod
    def is_supported(cls, filename: str) -> bool:
        ext = Path(filename).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def process_file(cls, filename: str, content_bytes: bytes) -> List[Dict[str, Any]]:
        """
        Processes document bytes and returns a list of page objects:
        [{"page": 1, "text": "extracted text content..."}]
        """
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            return cls._process_pdf(content_bytes)
        elif ext == ".docx":
            return cls._process_docx(content_bytes)
        elif ext in {".txt", ".md"}:
            return cls._process_text(content_bytes)
        else:
            raise ValueError(f"Unsupported file type: {ext}. Supported types: {cls.SUPPORTED_EXTENSIONS}")

    @staticmethod
    def _process_pdf(content_bytes: bytes) -> List[Dict[str, Any]]:
        stream = io.BytesIO(content_bytes)
        reader = pypdf.PdfReader(stream)
        pages = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            clean_text = text.strip()
            if clean_text:
                pages.append({
                    "page": idx + 1,
                    "text": clean_text
                })
        if not pages:
            # Fallback if no text extracted (e.g. empty pdf)
            pages.append({"page": 1, "text": ""})
        return pages

    @staticmethod
    def _process_docx(content_bytes: bytes) -> List[Dict[str, Any]]:
        stream = io.BytesIO(content_bytes)
        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs).strip()
        # In docx, exact physical page boundaries are layout-dependent;
        # We estimate pages by roughly 500 words per page.
        words = full_text.split()
        if not words:
            return [{"page": 1, "text": ""}]

        words_per_page = 500
        pages = []
        for i in range(0, len(words), words_per_page):
            page_text = " ".join(words[i:i + words_per_page])
            pages.append({
                "page": (i // words_per_page) + 1,
                "text": page_text
            })
        return pages

    @staticmethod
    def _process_text(content_bytes: bytes) -> List[Dict[str, Any]]:
        try:
            text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = content_bytes.decode("latin-1", errors="replace")

        # Estimate pages every ~3000 chars (~500 words)
        chunk_size = 3000
        clean_text = text.strip()
        if not clean_text:
            return [{"page": 1, "text": ""}]

        pages = []
        total_len = len(clean_text)
        for i in range(0, total_len, chunk_size):
            page_content = clean_text[i:i + chunk_size]
            pages.append({
                "page": (i // chunk_size) + 1,
                "text": page_content
            })
        return pages
